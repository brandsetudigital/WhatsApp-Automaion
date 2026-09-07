import os
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import qn, nsdecls

def set_cell_background(cell, fill_hex):
    tcPr = cell._element.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._element.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def create_hiring_knowledge_docx(filename="BrandSetu_Hiring_Knowledge_Base.docx"):
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.8)
        section.right_margin = Inches(0.8)

    # Styles & Colors
    PRIMARY_COLOR = RGBColor(14, 76, 146)    # Deep Navy Blue
    SECONDARY_COLOR = RGBColor(230, 81, 0)   # Vibrant Orange
    TEXT_DARK = RGBColor(33, 37, 41)         # Charcoal Dark
    GRAY_BG = "F4F6F9"
    HIGHLIGHT_BG = "FFF3E0"
    BLUE_BG = "EBF3FA"

    # Title
    title_p = doc.add_paragraph()
    title_p.paragraph_format.space_before = Pt(0)
    title_p.paragraph_format.space_after = Pt(4)
    title_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_title = title_p.add_run("🏢 BRAND SETU DIGITAL")
    run_title.font.name = "Arial"
    run_title.font.size = Pt(22)
    run_title.font.bold = True
    run_title.font.color.rgb = PRIMARY_COLOR

    # Subtitle
    sub_p = doc.add_paragraph()
    sub_p.paragraph_format.space_after = Pt(14)
    sub_p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run_sub = sub_p.add_run("Official WhatsApp Recruitment & Hiring Knowledge Base (Q&A Manual)")
    run_sub.font.name = "Arial"
    run_sub.font.size = Pt(13)
    run_sub.font.bold = True
    run_sub.font.color.rgb = SECONDARY_COLOR

    # Quick Info Box (Table)
    info_table = doc.add_table(rows=4, cols=2)
    info_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    info_data = [
        ("📍 Office Address", "103 Orange Business Park, Bhawarkua Main Road, Near Apple Hospital, Transport Nagar, Indore (M.P.) - 452014"),
        ("📞 Helpline / HR Contacts", "+91 9329232025 | +91 9244856276"),
        ("⏰ Working & Interview Hours", "Office: Mon–Sat (10:00 AM – 7:00 PM) | Interviews: Mon–Sat (10:00 AM – 6:00 PM) [Sunday Off]"),
        ("🏢 Work Mode & Types", "Strictly 100% In-Office (Indore) | Paid Internship (3-6 Mos) & Full-Time Career Opportunities")
    ]
    for i, (k, v) in enumerate(info_data):
        row = info_table.rows[i]
        c1, c2 = row.cells[0], row.cells[1]
        c1.width = Inches(2.2)
        c2.width = Inches(4.6)
        set_cell_background(c1, "EAEFF5")
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c1, 80, 80, 120, 120)
        set_cell_margins(c2, 80, 80, 120, 120)
        
        p1 = c1.paragraphs[0]
        r1 = p1.add_run(k)
        r1.font.bold = True
        r1.font.size = Pt(10)
        r1.font.color.rgb = PRIMARY_COLOR

        p2 = c2.paragraphs[0]
        r2 = p2.add_run(v)
        r2.font.size = Pt(9.5)
        r2.font.color.rgb = TEXT_DARK

    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    # Section 1: Active Roles
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r_h1 = h1.add_run("1. 🎯 ACTIVE JOB & INTERNSHIP OPENINGS (6 ROLES)")
    r_h1.font.name = "Arial"
    r_h1.font.size = Pt(14)
    r_h1.font.bold = True
    r_h1.font.color.rgb = PRIMARY_COLOR

    roles = [
        ("1. 🎬 Video Editor", "Reels, YouTube long-form, dynamic cuts, motion graphics, audio syncing, subtitle animations.", "Adobe Premiere Pro, After Effects, DaVinci Resolve, CapCut Pro", "Video samples / Google Drive link + PDF Resume"),
        ("2. 🤖 AI Video Expert", "AI video generation, prompt engineering, AI avatar creation, realistic animations for ads.", "Midjourney, Runway Gen-2/Gen-3, Kling AI, Luma Dream Machine, Pika, HeyGen", "AI Video creations / Drive link + PDF Resume"),
        ("3. 🎨 Graphic Designer", "Social media creatives, ad banners, typography, brand identity, thumbnail designs.", "Adobe Photoshop, Illustrator, Figma, Canva Pro", "Behance / Google Drive / Figma link + PDF Resume"),
        ("4. 🔎 SEO & AEO Expert", "On-Page, Technical SEO, AI Search Optimization (AEO for Perplexity, ChatGPT, Gemini), Page 1 rankings.", "Ahrefs, SEMrush, Google Search Console, Screaming Frog", "Live case studies / ranking proofs + PDF Resume"),
        ("5. 📱 Social Media Manager", "Organic profile growth, viral reels strategy, content calendar planning, caption writing.", "Meta Business Suite, Canva, Notion, Analytics tools", "Past managed profiles & growth case studies + PDF Resume"),
        ("6. 📢 Digital Marketing Manager", "High-ROI Paid Ads on Meta and Google, lead gen, media planning, performance analytics.", "Meta Ads Manager, Google Ads, GA4, Funnel Analytics", "Ad campaign case studies / ROAS track record + PDF Resume")
    ]

    role_table = doc.add_table(rows=1, cols=4)
    role_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    hdr_cells = role_table.rows[0].cells
    hdr_titles = ["Role Title", "Scope & Responsibilities", "Tools / Software", "Required Portfolio"]
    col_widths = [Inches(1.5), Inches(2.2), Inches(1.6), Inches(1.5)]

    for idx, name in enumerate(hdr_titles):
        cell = hdr_cells[idx]
        cell.width = col_widths[idx]
        set_cell_background(cell, "0E4C92")
        set_cell_margins(cell, 100, 100, 100, 100)
        p = cell.paragraphs[0]
        r = p.add_run(name)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    for role in roles:
        row = role_table.add_row()
        for idx, text in enumerate(role):
            cell = row.cells[idx]
            cell.width = col_widths[idx]
            set_cell_background(cell, GRAY_BG if role[0].startswith("1") or role[0].startswith("3") or role[0].startswith("5") else "FFFFFF")
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            r = p.add_run(text)
            r.font.size = Pt(9)
            if idx == 0:
                r.font.bold = True
                r.font.color.rgb = PRIMARY_COLOR

    doc.add_paragraph().paragraph_format.space_after = Pt(14)

    # Section 2: Complete Q&A
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(12)
    h2.paragraph_format.space_after = Pt(6)
    r_h2 = h2.add_run("2. 💬 CANDIDATE FREQUENTLY ASKED QUESTIONS & ANSWERS (Q&A)")
    r_h2.font.name = "Arial"
    r_h2.font.size = Pt(14)
    r_h2.font.bold = True
    r_h2.font.color.rgb = PRIMARY_COLOR

    qa_list = [
        {
            "q": "Q1. Salary / Package / Stipend kitna milega?",
            "topic": "Compensation Policy",
            "policy": "Salary/stipend is strictly evaluated based on experience, skill test, and in-person practical interview performance. It is finalized face-to-face during the interview.",
            "hinglish": "Hamare yahan salary / stipend aapke Experience, Skills aur In-Person Practical Interview ke basis par decide hoti hai aur interview ke dauraan finalize kar di jayegi. Kripya apna updated Resume (PDF) aur portfolio share karein taaki hum interview schedule kar sakein.",
            "english": "Salary and stipend are decided based on your Experience, Skills, and In-Person Practical Interview performance, and will be discussed and finalized during the interview. Please share your updated Resume (PDF) or portfolio link."
        },
        {
            "q": "Q2. Work From Home (WFH) / Remote available hai?",
            "topic": "Work Mode Policy",
            "policy": "100% In-Office at 103 Orange Business Park, Bhawarkua, Indore. No permanent remote work.",
            "hinglish": "Yeh position strictly Onsite In-Office role hai hamare Indore office (103 Orange Business Park, Bhawarkua) ke liye. Remote ya Work-From-Home option available nahi hai. Agar aap in-office comfortable hain toh Resume (PDF) share karein.",
            "english": "This is strictly an Onsite In-Office position at our Indore office (103 Orange Business Park, Bhawarkua). We currently do not offer remote or work-from-home options."
        },
        {
            "q": "Q3. Office Location & Landmark kya hai?",
            "topic": "Location & Directions",
            "policy": "103 Orange Business Park, Bhawarkua Main Road, Near Apple Hospital, Transport Nagar, Indore (M.P.) - 452014.",
            "hinglish": "Office Address: 103 Orange Business Park, Bhawarkua Main Road, Near Apple Hospital, Transport Nagar, Indore (M.P.) - 452014. Timing: Mon–Sat 10:00 AM – 7:00 PM. Helpline: +91 9329232025.",
            "english": "Office Location: 103 Orange Business Park, Bhawarkua Main Road, Near Apple Hospital, Transport Nagar, Indore (M.P.) - 452014. Office Timings: Mon–Sat 10:00 AM – 7:00 PM. Helpline: +91 9329232025."
        },
        {
            "q": "Q4. Office Timings & Interview Slots kya hain?",
            "topic": "Schedule & Working Hours",
            "policy": "Office: Mon–Sat 10:00 AM – 7:00 PM (Sunday Off). Interviews: Mon–Sat 10:00 AM – 6:00 PM.",
            "hinglish": "Office Working Hours: Monday to Saturday, 10:00 AM – 7:00 PM (Sunday Off). Interview Slots: Monday to Saturday, 10:00 AM – 6:00 PM. Aap inme se kisi bhi din slot schedule kar sakte hain.",
            "english": "Office Hours: Monday to Saturday, 10:00 AM – 7:00 PM (Sunday Off). Interview Slots: Monday to Saturday, 10:00 AM – 6:00 PM. You can schedule between 10:00 AM and 6:00 PM on any working day."
        },
        {
            "q": "Q5. Freshers apply kar sakte hain? Paid Internship details?",
            "topic": "Freshers & Internships",
            "policy": "Freshers get 3-6 months Paid Internship with monthly stipend, mentorship, certificate, and Full-Time PPO job offer.",
            "hinglish": "Haan, Freshers bilkul apply kar sakte hain! Freshers ke liye Paid Internship (3 to 6 Months) provide karte hain jisme: Handsome Stipend, Live Projects, Mentorship, Certificate aur Full-Time Job Offer (PPO) included hai.",
            "english": "Yes, Freshers are welcome! We offer Paid Internships (3 to 6 Months) that include: Monthly Stipend, Live Client Projects, Industry Mentorship, Certificate of Completion, and Full-Time Job Offer (PPO)."
        },
        {
            "q": "Q6. Resume & Portfolio format kya hona chahiye?",
            "topic": "Document Formats",
            "policy": "Strictly PDF (.pdf) for resumes. Photos/screenshots rejected. Drive/Behance/Figma links for portfolio.",
            "hinglish": "1. Resume strictly PDF format (.pdf) me share karein (Photos/camera screenshots accept nahi hote). 2. Portfolio ke liye apna Google Drive, Behance, ya Figma link yahan share karein.",
            "english": "1. Resumes must be in PDF format (.pdf) only (Photos or screenshots are not accepted). 2. For portfolio / work samples, please share your Google Drive, Behance, or Figma link."
        },
        {
            "q": "Q7. Selection & Interview Process kya hoga?",
            "topic": "Recruitment Stages",
            "policy": "WhatsApp Role Selection -> PDF Resume/Portfolio Review -> Indore In-Office Visit & Practical Skill Test -> Final HR Round -> Offer Letter & Joining.",
            "hinglish": "1. Resume & Portfolio Review on WhatsApp. 2. Indore Office Visit (103 Orange Business Park). 3. Practical Skill Assessment (Live task). 4. Final HR round -> Offer Letter & Immediate Joining!",
            "english": "1. WhatsApp Resume & Portfolio Review. 2. In-person visit at Indore Office. 3. Practical Skill Assessment (Live test). 4. Final HR Discussion -> Offer Letter & Onboarding!"
        },
        {
            "q": "Q8. Interview ke liye kya documents leke aana hai?",
            "topic": "Required Documents",
            "policy": "Updated Resume (Hard copy / PDF), Past Portfolio/work samples link (accessible on phone/laptop), and Aadhar ID proof.",
            "hinglish": "Interview ke liye: 1. Updated Resume (Hard copy / PDF), 2. Portfolio / Work samples (accessible on phone/laptop), 3. Government ID Proof (Aadhar Card).",
            "english": "Please bring: 1. Updated Resume (Hard copy / PDF), 2. Past Portfolio / Work samples (accessible on drive/laptop), 3. Government ID Proof (Aadhar Card)."
        },
        {
            "q": "Q9. Candidate Indore se bahar rehta hai (Online Round)?",
            "topic": "Outstation Candidates",
            "policy": "Initial round can be done via Google Meet online. Work is strictly in-office onsite at Indore upon joining.",
            "hinglish": "Agar aap abhi Indore se bahar hain, toh hum aapka Online Google Meet Interview conduct kar sakte hain (Mon–Sat, 10:00 AM – 6:00 PM). Joining ke baad job strictly Indore office onsite rahegi.",
            "english": "If you are currently outside Indore, we can conduct your initial round online via Google Meet (Mon–Sat, 10:00 AM – 6:00 PM). Upon selection, the job is strictly In-Office at Indore."
        },
        {
            "q": "Q10. Interview reschedule / time change kaise karein?",
            "topic": "Rescheduling Policy",
            "policy": "Flexible rescheduling within Mon–Sat (10:00 AM – 6:00 PM).",
            "hinglish": "Koi baat nahi! Aap apni suvidha ke anusaar Monday se Saturday (10:00 AM se 6:00 PM ke beech) koi bhi suitable Date aur Time bata dijiye, hum aapka slot update kar denge.",
            "english": "No problem at all! Please let us know your preferred Date and Time (Monday to Saturday, 10:00 AM – 6:00 PM) and we will update your interview schedule accordingly."
        }
    ]

    for item in qa_list:
        qa_box = doc.add_table(rows=4, cols=1)
        qa_box.alignment = WD_TABLE_ALIGNMENT.CENTER
        
        # Header Row (Question)
        c0 = qa_box.rows[0].cells[0]
        c0.width = Inches(6.8)
        set_cell_background(c0, "0E4C92")
        set_cell_margins(c0, 80, 80, 100, 100)
        p0 = c0.paragraphs[0]
        r0 = p0.add_run(f"❓ {item['q']}")
        r0.font.bold = True
        r0.font.size = Pt(11)
        r0.font.color.rgb = RGBColor(255, 255, 255)

        # Policy
        c1 = qa_box.rows[1].cells[0]
        c1.width = Inches(6.8)
        set_cell_background(c1, "FFF3E0")
        set_cell_margins(c1, 60, 60, 100, 100)
        p1 = c1.paragraphs[0]
        r1_title = p1.add_run("📌 Official Policy: ")
        r1_title.font.bold = True
        r1_title.font.size = Pt(9.5)
        r1_title.font.color.rgb = SECONDARY_COLOR
        r1_text = p1.add_run(item['policy'])
        r1_text.font.size = Pt(9.5)
        r1_text.font.color.rgb = TEXT_DARK

        # Hinglish Response
        c2 = qa_box.rows[2].cells[0]
        c2.width = Inches(6.8)
        set_cell_background(c2, "F8FAFC")
        set_cell_margins(c2, 60, 60, 100, 100)
        p2 = c2.paragraphs[0]
        r2_title = p2.add_run("💬 WhatsApp Reply (Hinglish): ")
        r2_title.font.bold = True
        r2_title.font.size = Pt(9.5)
        r2_title.font.color.rgb = PRIMARY_COLOR
        r2_text = p2.add_run(f'"{item["hinglish"]}"')
        r2_text.font.size = Pt(9)
        r2_text.font.italic = True
        r2_text.font.color.rgb = TEXT_DARK

        # English Response
        c3 = qa_box.rows[3].cells[0]
        c3.width = Inches(6.8)
        set_cell_background(c3, "FFFFFF")
        set_cell_margins(c3, 60, 60, 100, 100)
        p3 = c3.paragraphs[0]
        r3_title = p3.add_run("💬 WhatsApp Reply (English): ")
        r3_title.font.bold = True
        r3_title.font.size = Pt(9.5)
        r3_title.font.color.rgb = PRIMARY_COLOR
        r3_text = p3.add_run(f'"{item["english"]}"')
        r3_text.font.size = Pt(9)
        r3_text.font.italic = True
        r3_text.font.color.rgb = TEXT_DARK

        doc.add_paragraph().paragraph_format.space_after = Pt(8)

    # Section 3: HR WhatsApp Admin Commands
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(12)
    h3.paragraph_format.space_after = Pt(6)
    r_h3 = h3.add_run("3. 👮 HR WHATSAPP MOBILE ACTION COMMANDS")
    r_h3.font.name = "Arial"
    r_h3.font.size = Pt(14)
    r_h3.font.bold = True
    r_h3.font.color.rgb = PRIMARY_COLOR

    cmd_desc = doc.add_paragraph()
    r_cmd_desc = cmd_desc.add_run("HR admins can text these direct action commands from their mobile WhatsApp to update candidate statuses and trigger official messages automatically:")
    r_cmd_desc.font.size = Pt(9.5)

    cmd_table = doc.add_table(rows=1, cols=3)
    cmd_table.alignment = WD_TABLE_ALIGNMENT.CENTER
    c_hdr = cmd_table.rows[0].cells
    c_hdr[0].width, c_hdr[1].width, c_hdr[2].width = Inches(2.2), Inches(1.8), Inches(2.8)
    for idx, t in enumerate(["Command Example", "Action Executed", "Candidate Notification"]):
        cell = c_hdr[idx]
        set_cell_background(cell, "0E4C92")
        set_cell_margins(cell, 80, 80, 80, 80)
        p = cell.paragraphs[0]
        r = p.add_run(t)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    commands = [
        ("select 9876543210", "Marks Candidate 'Selected'", "Sends official Congratulations, Office Address & Offer letter details."),
        ("reject 9876543210", "Marks Candidate 'Rejected'", "Sends polite rejection & talent pool retention feedback."),
        ("hold 9876543210", "Marks Candidate 'On Hold'", "Sends update stating profile is under evaluation (decision in 2-3 days)."),
        ("status 9876543210", "Queries CRM Status", "Returns current role, interview date, resume status & notes to HR mobile.")
    ]

    for cmd, act, notif in commands:
        row = cmd_table.add_row()
        for idx, val in enumerate([cmd, act, notif]):
            cell = row.cells[idx]
            cell.width = [Inches(2.2), Inches(1.8), Inches(2.8)][idx]
            set_cell_background(cell, "F8FAFC" if idx % 2 == 0 else "FFFFFF")
            set_cell_margins(cell, 60, 60, 60, 60)
            p = cell.paragraphs[0]
            r = p.add_run(val)
            r.font.size = Pt(9)
            if idx == 0:
                r.font.bold = True
                r.font.color.rgb = SECONDARY_COLOR

    # Save document
    doc.save(filename)
    print("SUCCESS: Word Document generated: " + filename)

if __name__ == "__main__":
    create_hiring_knowledge_docx()
